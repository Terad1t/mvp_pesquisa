from __future__ import annotations

import json
import os
import queue
import threading
from pathlib import Path
from tkinter import filedialog

import customtkinter as ctk
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from dotenv import dotenv_values, load_dotenv, set_key

from extractor import ErroDeExtracao, ErroTransitorio
from main import _arquivos_pdf, _extrair_hibrido, _mover_pdf_processado, processar

CARGOS = ("GOVERNADOR", "SENADOR", "PRESIDENTE")
LABELS = {"GOVERNADOR": "Governador", "SENADOR": "Senador", "PRESIDENTE": "Presidente"}
ROOT = Path(__file__).parent


def export_json_to_xlsx(json_path: Path, xlsx_path: Path) -> None:
    """Converte o contrato JSON em abas de tabelas prontas para leitura e filtro."""
    data = json.loads(json_path.read_text(encoding="utf-8"))
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"

    navy = "123247"
    teal = "0B9DAA"
    pale = "E9EEF1"
    grid = Side(style="thin", color="D3DDE2")
    summary.append(["Campo", "Valor"])
    summary.append(["Estado", data.get("estado", "")])
    summary.append(["Completo", "Sim" if data.get("completo") else "Não"])
    summary.append(["Avisos", " | ".join(data.get("avisos", []))])
    summary.append(["Cargos pendentes", ", ".join(data.get("cargos_pendentes", []))])
    summary.append(["Cargos faltando", ", ".join(data.get("cargos_faltando", []))])
    _style_summary(summary, navy, teal, grid)

    cargos = data.get("cargos", {})
    if isinstance(cargos, list):
        cargos = {item.get("cargo", "CARGO"): item for item in cargos}
    for cargo, bloco in cargos.items():
        sheet = workbook.create_sheet(LABELS.get(cargo, cargo)[:31])
        _write_cargo_sheet(sheet, data.get("estado", ""), cargo, bloco, navy, teal, pale, grid)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)


def _style_summary(sheet, navy: str, teal: str, grid: Side) -> None:
    """Aplica o mesmo acabamento visual básico à aba de metadados."""
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=grid)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 80
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{sheet.max_row}"
    sheet.sheet_view.showGridLines = False


def _write_cargo_sheet(
    sheet,
    estado: str,
    cargo: str,
    bloco: dict,
    navy: str,
    teal: str,
    pale: str,
    grid: Side,
) -> None:
    headers = ["Ranking", "Candidato", "Partido", "Votos válidos", "% total"]
    sheet.merge_cells("A1:E1")
    sheet["A1"] = f"{estado} - {LABELS.get(cargo, cargo).upper()}"
    sheet["A1"].font = Font(name="Arial", size=18, bold=True, color=navy)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells("A2:E2")
    sheet["A2"] = _subtitle(cargo)
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="637682")
    sheet["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 26

    sheet.merge_cells("A4:E4")
    sheet["A4"] = f"{_section_number(cargo)}. {LABELS.get(cargo, cargo).upper()}"
    sheet["A4"].font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    sheet["A4"].fill = PatternFill("solid", fgColor=navy)
    sheet["A4"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[4].height = 25

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column, value=header)
        cell.font = Font(name="Arial", size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=teal)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[5].height = 22

    candidatos = sorted(
        bloco.get("candidatos", []),
        key=lambda item: (item.get("posicao") is None, item.get("posicao") or 0),
    )
    for candidate in candidatos:
        sheet.append([
            candidate.get("posicao"),
            candidate.get("nome"),
            candidate.get("partido"),
            candidate.get("porcentagem_valida"),
            candidate.get("porcentual"),
        ])

    sheet.append([None, "OUTROS", None, bloco.get("outros_valido"), bloco.get("outros_total")])
    sheet.append([None, "NS/NR", None, None, bloco.get("ns_nr")])
    sheet.append([None, "BRANCO/NULO", None, None, bloco.get("brancos_nulos")])

    for row in sheet.iter_rows(min_row=6, max_row=sheet.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = Border(bottom=grid)
            cell.alignment = Alignment(vertical="center")
        if row[1].value in {"OUTROS", "NS/NR", "BRANCO/NULO"}:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=pale)
    for row in range(6, sheet.max_row + 1):
        for column in (4, 5):
            sheet.cell(row=row, column=column).number_format = '0.0"%"'
    sheet.auto_filter.ref = f"A5:E{sheet.max_row}"
    sheet.freeze_panes = "A6"
    sheet.sheet_view.showGridLines = False
    for column, width in {"A": 11, "B": 31, "C": 20, "D": 17, "E": 13}.items():
        sheet.column_dimensions[column].width = width
    sheet.print_title_rows = "1:5"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _subtitle(cargo: str) -> str:
    if cargo == "SENADOR":
        return "Os votos foram tratados separadamente. O consolidado utiliza exclusivamente a porcentagem de casos da tabela oficial."
    return "Percentuais organizados conforme a tabela oficial da pesquisa."


def _section_number(cargo: str) -> int:
    return {"GOVERNADOR": 1, "SENADOR": 3, "PRESIDENTE": 5}.get(cargo, 1)


class DesktopApp(ctk.CTk):
    def __init__(self) -> None:
        super().__init__()
        self.title("MVP Pesquisa | Veritá → PesquisaPRO")
        self.geometry("1120x760")
        self.minsize(900, 620)
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("green")
        self.events: queue.Queue[tuple[str, object]] = queue.Queue()
        self.worker: threading.Thread | None = None
        self.pdfs: list[Path] = []
        self.cards: dict[str, ctk.CTkLabel] = {}
        self.pdf_folder = str(ROOT / "PDFs")
        self._show_shell("processar")
        self.after(100, self._consume_events)

    def _clear_content(self) -> None:
        if hasattr(self, "content"):
            for child in self.content.winfo_children():
                child.destroy()

    def _show_shell(self, page: str) -> None:
        for child in self.winfo_children():
            child.destroy()
        self.sidebar = ctk.CTkFrame(self, width=215, corner_radius=0, fg_color="#15221f")
        self.sidebar.pack(side="left", fill="y")
        ctk.CTkLabel(self.sidebar, text="P", text_color="#c4ef62", font=ctk.CTkFont(size=34, weight="bold")).pack(anchor="w", padx=28, pady=(30, 0))
        ctk.CTkLabel(self.sidebar, text="PESQUISAS", text_color="#8fa49d", font=ctk.CTkFont(size=10, weight="bold")).pack(anchor="w", padx=30, pady=(0, 42))
        self._nav_button("Processar", "processar", page)
        self._nav_button("Dados / XLSX", "dados", page)
        self._nav_button("Bloqueados", "bloqueados", page)
        self._nav_button("Histórico", "historico", page)
        self._nav_button("Configurações", "configuracoes", page)
        ctk.CTkLabel(self.sidebar, text="parser → validação\n→ cálculo", text_color="#60756d", justify="left").pack(side="bottom", anchor="w", padx=30, pady=28)
        self.content = ctk.CTkFrame(self, fg_color="#101a19", corner_radius=0)
        self.content.pack(side="left", fill="both", expand=True)
        {"processar": self._show_processar, "dados": self._show_dados, "bloqueados": self._show_bloqueados, "historico": self._show_historico, "configuracoes": self._show_configuracoes}[page]()

    def _nav_button(self, text: str, page: str, active: str) -> None:
        ctk.CTkButton(self.sidebar, text=text, anchor="w", height=40, corner_radius=8, fg_color="#2b493e" if page == active else "transparent", hover_color="#29473d", command=lambda: self._show_shell(page)).pack(fill="x", padx=16, pady=3)

    def _header(self, title: str, subtitle: str) -> None:
        head = ctk.CTkFrame(self.content, fg_color="transparent")
        head.pack(fill="x", padx=40, pady=(34, 24))
        ctk.CTkLabel(head, text=title, font=ctk.CTkFont(size=30, weight="bold")).pack(anchor="w")
        ctk.CTkLabel(head, text=subtitle, text_color="#8fa49d", font=ctk.CTkFont(size=13)).pack(anchor="w", pady=(5, 0))

    def _show_processar(self) -> None:
        self._header("Processar relatórios", "Escolha um PDF ou uma pasta. O backend continuará responsivo durante retries do Gemini.")
        body = ctk.CTkFrame(self.content, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=40)
        picker = ctk.CTkFrame(body, corner_radius=14)
        picker.pack(fill="x", pady=(0, 18))
        self.selection = ctk.CTkLabel(picker, text="Nenhum relatório selecionado", anchor="w", text_color="#9aa9a3")
        self.selection.pack(fill="x", padx=20, pady=(20, 12))
        buttons = ctk.CTkFrame(picker, fg_color="transparent")
        buttons.pack(fill="x", padx=20, pady=(0, 20))
        ctk.CTkButton(buttons, text="Selecionar PDF", command=self._select_pdf, width=160).pack(side="left", padx=(0, 10))
        ctk.CTkButton(buttons, text="Selecionar pasta", command=self._select_folder, width=160, fg_color="#29473d", hover_color="#356050").pack(side="left")
        self.start_button = ctk.CTkButton(body, text="Iniciar processamento", command=self._start_processing, state="disabled", height=42, font=ctk.CTkFont(weight="bold"))
        self.start_button.pack(anchor="e", pady=(10, 25))
        self.run_area = ctk.CTkFrame(body, fg_color="transparent")
        self.run_area.pack(fill="both", expand=True)
        self._build_status_cards()

    def _build_status_cards(self) -> None:
        for cargo in CARGOS:
            card = ctk.CTkFrame(self.run_area, corner_radius=11)
            card.pack(fill="x", pady=5)
            ctk.CTkLabel(card, text=LABELS[cargo], font=ctk.CTkFont(size=15, weight="bold"), anchor="w").pack(side="left", padx=18, pady=15)
            label = ctk.CTkLabel(card, text="Aguardando", text_color="#8fa49d")
            label.pack(side="right", padx=18, pady=15)
            self.cards[cargo] = label
        self.run_status = ctk.CTkLabel(self.run_area, text="", text_color="#8fa49d", anchor="w")
        self.run_status.pack(fill="x", pady=(14, 0))

    def _select_pdf(self) -> None:
        selected = filedialog.askopenfilename(title="Selecionar relatório", filetypes=[("PDF", "*.pdf")])
        if selected:
            self.pdfs = [Path(selected)]
            self.selection.configure(text=self.pdfs[0].name, text_color="#edf3ef")
            self.start_button.configure(state="normal")

    def _select_folder(self) -> None:
        selected = filedialog.askdirectory(title="Selecionar pasta com PDFs")
        if not selected:
            return
        try:
            self.pdfs = _arquivos_pdf(Path(selected))
        except ValueError as error:
            self.selection.configure(text=str(error), text_color="#f0a58c")
            return
        self.pdf_folder = selected
        self.selection.configure(text=f"{len(self.pdfs)} PDF(s) em {Path(selected).name}", text_color="#edf3ef")
        self.start_button.configure(state="normal" if self.pdfs else "disabled")

    def _start_processing(self, paths: list[Path] | None = None) -> None:
        paths = paths or list(self.pdfs)
        if not paths or self.worker and self.worker.is_alive():
            return
        self._show_shell("processar")
        self.pdfs = paths
        self.worker = threading.Thread(target=self._worker, args=(paths,), daemon=True)
        self.worker.start()

    def _worker(self, paths: list[Path]) -> None:
        load_dotenv()
        for index, path in enumerate(paths, 1):
            self.events.put(("run", f"PDF {index}/{len(paths)}: {path.name} | parser/Gemini em execução"))
            try:
                result = _extrair_hibrido(path, None)
                output_name = "pesquisa.json" if len(paths) == 1 else f"{path.stem}.json"
                def on_cargo(cargo: str, status: str, issues: list[object]) -> None:
                    self.events.put(("cargo", (cargo, status, issues)))
                code = processar(result, nome_saida=output_name, on_cargo=on_cargo)
                _mover_pdf_processado(path, code)
                self.events.put(("run", f"{path.name} concluído (código {code})"))
            except (ErroTransitorio, ErroDeExtracao) as error:
                self.events.put(("error", f"{path.name}: {error}"))
            except Exception as error:
                self.events.put(("error", f"{path.name}: erro inesperado: {error}"))
        self.events.put(("done", None))

    def _consume_events(self) -> None:
        try:
            while True:
                kind, payload = self.events.get_nowait()
                if kind == "run" and hasattr(self, "run_status"):
                    self.run_status.configure(text=str(payload))
                elif kind == "cargo" and hasattr(self, "cards"):
                    cargo, status, _issues = payload  # type: ignore[misc]
                    if cargo in self.cards:
                        labels = {"ok": "✓ OK", "bloqueado": "⚠ Bloqueado", "pulado": "ℹ Pendente", "processando": "⏳ Processando..."}
                        colors = {"ok": "#b9ee78", "bloqueado": "#f0a58c", "pulado": "#f0cb79", "processando": "#8fc9ab"}
                        self.cards[cargo].configure(text=labels.get(status, status), text_color=colors.get(status, "#9aa9a3"))
                elif kind == "error" and hasattr(self, "run_status"):
                    self.run_status.configure(text=str(payload), text_color="#f0a58c")
                elif kind == "done" and hasattr(self, "run_status"):
                    self.run_status.configure(text="Processamento finalizado.", text_color="#b9ee78")
                elif kind == "export_done" and hasattr(self, "export_status"):
                    self.export_status.configure(text=f"XLSX salvo em {payload}", text_color="#b9ee78")
                    self.export_button.configure(state="normal")
                elif kind == "export_error" and hasattr(self, "export_status"):
                    self.export_status.configure(text=str(payload), text_color="#f0a58c")
                    self.export_button.configure(state="normal")
        except queue.Empty:
            pass
        self.after(100, self._consume_events)

    def _show_dados(self) -> None:
        self._header("Dados e exportação", "Consulte as tabelas estruturadas do JSON e gere uma planilha organizada.")
        reports = sorted((ROOT / "saidas").glob("*/pesquisa*.json")) if (ROOT / "saidas").is_dir() else []
        if not reports:
            ctk.CTkLabel(self.content, text="Nenhum JSON encontrado em saidas/.", text_color="#8fa49d").pack(anchor="w", padx=40, pady=30)
            return

        toolbar = ctk.CTkFrame(self.content, fg_color="transparent")
        toolbar.pack(fill="x", padx=40, pady=(0, 14))
        options = {str(index): path for index, path in enumerate(reports)}
        self.data_report_options = options
        self.data_report_var = ctk.StringVar(value=next(iter(options)))
        ctk.CTkLabel(toolbar, text="Relatório", text_color="#8fa49d").pack(side="left", padx=(0, 8))
        ctk.CTkOptionMenu(toolbar, variable=self.data_report_var, values=list(options), command=lambda _value: self._render_json_table()).pack(side="left")
        self.export_button = ctk.CTkButton(toolbar, text="Gerar XLSX", width=135, command=self._export_selected_json)
        self.export_button.pack(side="right")
        self.export_status = ctk.CTkLabel(self.content, text="", text_color="#8fa49d", anchor="e")
        self.export_status.pack(fill="x", padx=40)
        self.data_summary = ctk.CTkFrame(self.content, fg_color="#182823", corner_radius=10)
        self.data_summary.pack(fill="x", padx=40, pady=(10, 0))
        self.data_table = ctk.CTkScrollableFrame(self.content, fg_color="transparent")
        self.data_table.pack(fill="both", expand=True, padx=32, pady=(10, 24))
        self._render_json_table()

    def _render_json_table(self) -> None:
        if not hasattr(self, "data_table"):
            return
        for child in self.data_table.winfo_children():
            child.destroy()
        report = self.data_report_options[self.data_report_var.get()]
        try:
            data = json.loads(report.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as error:
            ctk.CTkLabel(self.data_table, text=f"Não foi possível ler o JSON: {error}", text_color="#f0a58c").pack(anchor="w")
            return
        for child in self.data_summary.winfo_children():
            child.destroy()
        estado = data.get("estado", report.parent.name)
        status = "Completo" if data.get("completo") else "Revisão necessária"
        status_color = "#b9ee78" if data.get("completo") else "#f0a58c"
        ctk.CTkLabel(self.data_summary, text=estado, font=ctk.CTkFont(size=18, weight="bold"), anchor="w").pack(side="left", padx=16, pady=12)
        ctk.CTkLabel(self.data_summary, text=status, text_color=status_color, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=12)
        ctk.CTkLabel(self.data_summary, text=f"Pendentes: {', '.join(data.get('cargos_pendentes', [])) or 'nenhum'}  |  Faltando: {', '.join(data.get('cargos_faltando', [])) or 'nenhum'}", text_color="#8fa49d", anchor="e").pack(side="right", padx=16, pady=12)
        if data.get("avisos"):
            ctk.CTkLabel(self.data_table, text="Avisos: " + " | ".join(data["avisos"]), text_color="#f0cb79", anchor="w", wraplength=780, justify="left").pack(fill="x", padx=8, pady=(8, 10))
        for cargo, bloco in data.get("cargos", {}).items():
            cargo_card = ctk.CTkFrame(self.data_table, corner_radius=12, fg_color="#111f1b")
            cargo_card.pack(fill="x", padx=4, pady=(12, 10))
            cargo_head = ctk.CTkFrame(cargo_card, fg_color="transparent")
            cargo_head.pack(fill="x", padx=14, pady=(12, 5))
            ctk.CTkLabel(cargo_head, text=LABELS.get(cargo, cargo), text_color="#c4ef62", font=ctk.CTkFont(size=18, weight="bold"), anchor="w").pack(side="left")
            ctk.CTkLabel(cargo_head, text=f"Origem: {bloco.get('origem', 'desconhecida')}", text_color="#8fa49d", anchor="e").pack(side="right")
            metrics = ctk.CTkFrame(cargo_card, fg_color="transparent")
            metrics.pack(fill="x", padx=14, pady=(0, 9))
            for label, value in (("NS/NR", bloco.get("ns_nr", "—")), ("Brancos/nulos", bloco.get("brancos_nulos", "—")), ("Outros válidos", bloco.get("outros_valido", "—")), ("Outros total", bloco.get("outros_total", "—"))):
                ctk.CTkLabel(metrics, text=f"{label}: {value}", text_color="#9eb5aa", anchor="w").pack(side="left", padx=(0, 24))
            header = ctk.CTkFrame(cargo_card, fg_color="#29473d")
            header.pack(fill="x", padx=8)
            columns = ("Posição", "Nome", "Partido", "Votos", "% total", "% válido")
            for column in columns:
                ctk.CTkLabel(header, text=column, anchor="w", font=ctk.CTkFont(size=11, weight="bold")).pack(side="left", expand=True, fill="x", padx=8, pady=8)
            candidatos = sorted(bloco.get("candidatos", []), key=lambda item: (item.get("posicao") is None, item.get("posicao") or 0))
            for candidate in candidatos:
                row = ctk.CTkFrame(cargo_card, fg_color="#182823")
                row.pack(fill="x", padx=8, pady=1)
                values = (candidate.get("posicao", "—"), candidate.get("nome", "—"), candidate.get("partido", "—"), candidate.get("votos", "—"), candidate.get("porcentual", "—"), candidate.get("porcentagem_valida", "—"))
                for value in values:
                    ctk.CTkLabel(row, text=str(value), anchor="w", font=ctk.CTkFont(size=11)).pack(side="left", expand=True, fill="x", padx=8, pady=7)

    def _export_selected_json(self) -> None:
        report = self.data_report_options[self.data_report_var.get()]
        target = filedialog.asksaveasfilename(title="Salvar planilha", initialfile=f"{report.stem}.xlsx", defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not target:
            return
        self.export_button.configure(state="disabled")
        self.export_status.configure(text="Gerando planilha...", text_color="#8fc9ab")
        threading.Thread(target=self._export_worker, args=(report, Path(target)), daemon=True).start()

    def _export_worker(self, report: Path, target: Path) -> None:
        try:
            export_json_to_xlsx(report, target)
            self.events.put(("export_done", target))
        except Exception as error:
            self.events.put(("export_error", f"Falha ao gerar XLSX: {error}"))

    def _show_bloqueados(self) -> None:
        self._header("Fila de revisão", "PDFs bloqueados pelo validator ficam aqui para reprocessamento. Os números não são editados pela interface.")
        folder = ROOT / "PDFs" / "bloqueados"
        files = sorted(folder.glob("*.pdf")) if folder.is_dir() else []
        if not files:
            ctk.CTkLabel(self.content, text="Nenhum PDF bloqueado encontrado.", text_color="#8fa49d").pack(anchor="w", padx=40, pady=30)
            return
        scroll = ctk.CTkScrollableFrame(self.content, fg_color="transparent")
        scroll.pack(fill="both", expand=True, padx=32, pady=(0, 25))
        for pdf in files:
            card = ctk.CTkFrame(scroll, corner_radius=12)
            card.pack(fill="x", pady=6)
            ctk.CTkLabel(card, text=pdf.name, font=ctk.CTkFont(weight="bold"), anchor="w").pack(side="left", padx=18, pady=16)
            ctk.CTkLabel(card, text="Revisão necessária", text_color="#f0a58c").pack(side="left", padx=12)
            ctk.CTkButton(card, text="Reprocessar", width=120, command=lambda item=pdf: self._start_processing([item])).pack(side="right", padx=18, pady=10)

    def _show_historico(self) -> None:
        self._header("Histórico por estado", "Consulte os JSONs produzidos sem abrir os arquivos manualmente.")
        scroll = ctk.CTkScrollableFrame(self.content, fg_color="transparent")
        scroll.pack(fill="both", expand=True, padx=32, pady=(0, 25))
        root = ROOT / "saidas"
        reports = sorted(root.glob("*/pesquisa*.json")) if root.is_dir() else []
        if not reports:
            ctk.CTkLabel(scroll, text="Nenhum relatório gerado ainda.", text_color="#8fa49d").pack(anchor="w", padx=8, pady=20)
            return
        for report in reports:
            try:
                data = json.loads(report.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                continue
            card = ctk.CTkFrame(scroll, corner_radius=12)
            card.pack(fill="x", pady=6)
            left = ctk.CTkFrame(card, fg_color="transparent")
            left.pack(side="left", fill="x", expand=True, padx=18, pady=14)
            ctk.CTkLabel(left, text=data.get("estado", report.parent.name), font=ctk.CTkFont(size=16, weight="bold"), anchor="w").pack(anchor="w")
            ctk.CTkLabel(left, text=report.name, text_color="#8fa49d", anchor="w").pack(anchor="w", pady=(3, 0))
            status = "Completo" if data.get("completo") else "Incompleto"
            color = "#b9ee78" if data.get("completo") else "#f0a58c"
            right = ctk.CTkFrame(card, fg_color="transparent")
            right.pack(side="right", padx=18, pady=14)
            ctk.CTkLabel(right, text=status, text_color=color, font=ctk.CTkFont(weight="bold")).pack(anchor="e")
            pending = ", ".join(data.get("cargos_pendentes", [])) or "nenhum"
            missing = ", ".join(data.get("cargos_faltando", [])) or "nenhum"
            ctk.CTkLabel(right, text=f"Pendentes: {pending} | Faltando: {missing}", text_color="#8fa49d", font=ctk.CTkFont(size=11)).pack(anchor="e", pady=(4, 0))

    def _show_configuracoes(self) -> None:
        self._header("Configurações", "Ajuste o caminho de entrada e a chave usada pelo fallback Gemini.")
        panel = ctk.CTkFrame(self.content, corner_radius=14)
        panel.pack(fill="x", padx=40, pady=(0, 20))
        env_path = ROOT / ".env"
        values = dotenv_values(env_path)
        ctk.CTkLabel(panel, text="Pasta padrão de PDFs", anchor="w").pack(fill="x", padx=22, pady=(22, 6))
        folder_var = ctk.StringVar(value=self.pdf_folder)
        folder_entry = ctk.CTkEntry(panel, textvariable=folder_var)
        folder_entry.pack(fill="x", padx=22, pady=(0, 15))
        ctk.CTkButton(panel, text="Escolher pasta", command=lambda: self._choose_config_folder(folder_var), width=140).pack(anchor="w", padx=22, pady=(0, 18))
        ctk.CTkLabel(panel, text="Chave GEMINI_API_KEY", anchor="w").pack(fill="x", padx=22, pady=(0, 6))
        key_var = ctk.StringVar(value=values.get("GEMINI_API_KEY", ""))
        ctk.CTkEntry(panel, textvariable=key_var, show="•").pack(fill="x", padx=22, pady=(0, 20))
        message = ctk.CTkLabel(panel, text="", text_color="#b9ee78")
        message.pack(anchor="w", padx=22)
        def save() -> None:
            self.pdf_folder = folder_var.get().strip()
            env_path.touch(exist_ok=True)
            set_key(str(env_path), "GEMINI_API_KEY", key_var.get().strip())
            message.configure(text="Configurações salvas em .env")
        ctk.CTkButton(panel, text="Salvar configurações", command=save, height=38).pack(anchor="e", padx=22, pady=(8, 22))
        ctk.CTkLabel(self.content, text="A chave é usada somente pelo backend Gemini. Os números extraídos não são editáveis nesta tela.", text_color="#8fa49d", wraplength=700, justify="left").pack(anchor="w", padx=40)

    def _choose_config_folder(self, variable: ctk.StringVar) -> None:
        selected = filedialog.askdirectory(title="Pasta padrão de PDFs")
        if selected:
            variable.set(selected)


if __name__ == "__main__":
    DesktopApp().mainloop()
